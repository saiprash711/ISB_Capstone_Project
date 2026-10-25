"""
Generate Comprehensive NotebookLM PPT Presentation Script & Content Dossier in Word (.docx)
Project: Demand Prediction by SKU - Sales Data Collection and EDA (ISB Capstone / Deloitte)
Authors: Sai Prashanth Sarathy, R Arvindshankar, Hariharan N, Vedika Sharma, Vani Som
"""

import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Colors
HEX_NAVY = "0E1622"
HEX_CYAN = "0284C7"
HEX_TEAL = "00A8CC"
HEX_DELOITTE_GREEN = "86BC25"
HEX_DARK_GRAY = "334155"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_WHITE = "FFFFFF"
HEX_CARD_BG = "F1F5F9"
HEX_AMBER = "D97706"
HEX_HEADER_BG = "0F172A"

COLOR_NAVY = RGBColor(14, 22, 34)
COLOR_CYAN = RGBColor(2, 132, 199)
COLOR_TEAL = RGBColor(0, 168, 204)
COLOR_DELOITTE = RGBColor(134, 188, 37)
COLOR_DARK = RGBColor(51, 65, 85)
COLOR_MUTED = RGBColor(100, 116, 139)
COLOR_WHITE = RGBColor(255, 255, 255)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = ['<w:tcBorders ' + nsdecls('w') + '>']
    for side, border in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        if border:
            borders.append(f'<w:{side} w:val="{border.get("val", "single")}" w:sz="{border.get("sz", "4")}" w:space="0" w:color="{border.get("color", "CBD5E1")}"/>')
        else:
            borders.append(f'<w:{side} w:val="none"/>')
    borders.append('</w:tcBorders>')
    tcPr.append(parse_xml(''.join(borders)))

def add_header_footer(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run("ISB AMPBA Capstone  |  Deloitte  |  Enterprise Demand Intelligence Platform")
    hrun.font.name = "Calibri"
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = COLOR_MUTED
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run("NotebookLM Master Source & Presentation Script Dossier  •  Page ")
    frun.font.name = "Calibri"
    frun.font.size = Pt(8.5)
    frun.font.color.rgb = COLOR_MUTED
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    fp._p.append(fldSimple)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = COLOR_NAVY
    
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(0)
    p_line.paragraph_format.space_after = Pt(8)
    p_line.paragraph_format.keep_with_next = True
    run_bar = p_line.add_run("―" * 58)
    run_bar.font.name = 'Calibri'
    run_bar.font.size = Pt(8.5)
    run_bar.font.color.rgb = COLOR_CYAN
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(13)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(12.5)
    run.bold = True
    run.font.color.rgb = COLOR_CYAN
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = 'Calibri'
    run.font.size = Pt(10.5)
    run.bold = True
    run.font.color.rgb = COLOR_DARK
    return p

def add_body_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(9.5)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(9.5)
    r_body.italic = italic
    r_body.font.color.rgb = COLOR_DARK
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(9.5)
        r_pre.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(9.5)
    r_body.font.color.rgb = COLOR_DARK
    return p

def add_callout(doc, title, paragraphs, border_color=HEX_CYAN, bg_color=HEX_CARD_BG, icon="📌"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.9)
    
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    border_xml = (
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="20" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(parse_xml(border_xml))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    run_title = p.add_run(f"{icon} {title.upper()}")
    run_title.bold = True
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(10)
    run_title.font.color.rgb = COLOR_CYAN if border_color == HEX_CYAN else (COLOR_DELOITTE if border_color == HEX_DELOITTE_GREEN else COLOR_DARK)
    
    for item in paragraphs:
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(1.5)
        p2.paragraph_format.space_after = Pt(3)
        p2.paragraph_format.line_spacing = 1.15
        run_text = p2.add_run(item)
        run_text.font.name = 'Calibri'
        run_text.font.size = Pt(9.0)
        run_text.font.color.rgb = COLOR_DARK
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(9.0)
            run.font.bold = True
            run.font.color.rgb = COLOR_WHITE
            
    # Data rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = HEX_WHITE if r_idx % 2 == 0 else HEX_CARD_BG
        for c_idx, cell_value in enumerate(row_data):
            row_cells[c_idx].text = str(cell_value)
            set_cell_background(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.name = 'Calibri'
                run.font.size = Pt(8.5)
                run.font.color.rgb = COLOR_DARK
                
    for row in table.rows:
        for c_idx, cell in enumerate(row.cells):
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            set_cell_borders(
                cell,
                top={'sz': 2, 'color': HEX_BORDER},
                bottom={'sz': 2, 'color': HEX_BORDER},
                left=None, right=None
            )
            
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)

print("Core UI elements ready. Writing complete document...")
