"""
Helpers that write content into the ISB Capstone Report Template while reusing
its own styles (Normal, Heading 1-3, Caption, Source Note, Table Head/Text,
Box Title/Text, Reference Entry) and its live SEQ caption fields.
"""
import copy
import re
from xml.sax.saxutils import escape

from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NSDECL = (f'xmlns:w="{W_NS}" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')

BULLET_NUM = 1        # numbering instance used for body bullets (● list)
TAKEAWAY_NUM = 3      # numbering instance used inside Key Takeaways boxes
CELL_MAR = ('<w:tcMar><w:top w:w="80" w:type="dxa"/><w:left w:w="120" w:type="dxa"/>'
            '<w:bottom w:w="80" w:type="dxa"/><w:right w:w="120" w:type="dxa"/></w:tcMar>')
TBL_PR = ('<w:tblPr><w:tblW w:w="9360" w:type="dxa"/><w:tblBorders>'
          + "".join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="BFC8D6"/>'
                    for s in ("top", "left", "bottom", "right", "insideH", "insideV"))
          + '</w:tblBorders><w:tblLayout w:type="fixed"/><w:tblCellMar><w:left w:w="10" w:type="dxa"/>'
          '<w:right w:w="10" w:type="dxa"/></w:tblCellMar><w:tblLook w:val="04A0" w:firstRow="1" '
          'w:lastRow="0" w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/></w:tblPr>')

_TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*|\^\^.+?\^\^)")


def runs_xml(text, base_rpr=""):
    """Inline markup: **bold**, *italic*, ^^superscript^^."""
    out = []
    for part in _TOKEN.split(text):
        if not part:
            continue
        rpr = base_rpr
        if part.startswith("**") and part.endswith("**"):
            part, rpr = part[2:-2], base_rpr + "<w:b/><w:bCs/>"
        elif part.startswith("^^") and part.endswith("^^"):
            part, rpr = part[2:-2], base_rpr + '<w:vertAlign w:val="superscript"/>'
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, rpr = part[1:-1], base_rpr + "<w:i/><w:iCs/>"
        rp = f"<w:rPr>{rpr}</w:rPr>" if rpr else ""
        out.append(f'<w:r>{rp}<w:t xml:space="preserve">{escape(part)}</w:t></w:r>')
    return "".join(out)


def p_xml(text, style=None, ppr_extra="", rpr=""):
    st = f'<w:pStyle w:val="{style}"/>' if style else ""
    ppr = f"<w:pPr>{st}{ppr_extra}</w:pPr>" if (st or ppr_extra) else ""
    return f"<w:p {NSDECL}>{ppr}{runs_xml(text, rpr)}</w:p>"


class Writer:
    """Appends elements after a moving cursor inside the document body."""

    def __init__(self, doc, cursor, protos):
        self.doc = doc
        self.cursor = cursor
        self.protos = protos
        self.fig_no = {}
        self.tab_no = {}
        self.chapter = 0
        self.next_num_id = 20
        self.word_count = 0
        self.pending_reset = set()

    def section(self, label):
        """Start a caption-numbering section without a Heading 1 (e.g., Appendix B):
        the next Figure/Table caption restarts at 1 via the SEQ field's \\r switch."""
        self.chapter = label
        self.pending_reset = {"Figure", "Table"}

    # -- low level ---------------------------------------------------------
    def _add(self, el):
        self.cursor.addnext(el)
        self.cursor = el
        return el

    def raw(self, xml):
        return self._add(parse_xml(xml))

    def _count(self, text):
        self.word_count += len(re.findall(r"[A-Za-z0-9][\w'’%.,-]*", text))

    # -- text --------------------------------------------------------------
    def h1(self, text, chapter=None):
        if chapter is not None:
            self.chapter = chapter
        return self.raw(p_xml(text, "Heading1"))

    def h2(self, text):
        return self.raw(p_xml(text, "Heading2"))

    def h3(self, text):
        return self.raw(p_xml(text, "Heading3"))

    def p(self, text, style=None, ppr_extra="", count=True):
        if count:
            self._count(text)
        return self.raw(p_xml(text, style, ppr_extra))

    def paras(self, block):
        for para in [x.strip() for x in block.strip().split("\n\n") if x.strip()]:
            self.p(" ".join(para.split()))

    def bullets(self, items):
        for it in items:
            self._count(it)
            self.raw(p_xml(it, "ListParagraph",
                           f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{BULLET_NUM}"/></w:numPr>'
                           '<w:spacing w:after="80"/>'))

    def numbered(self, items):
        num_id = self._new_numbering()
        for it in items:
            self._count(it)
            self.raw(p_xml(it, "ListParagraph",
                           f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'
                           '<w:spacing w:after="80"/>'))

    def _new_numbering(self):
        """A fresh decimal list (abstractNum 0) that restarts at 1."""
        numbering = self.doc.part.numbering_part.element
        num_id = self.next_num_id
        self.next_num_id += 1
        numbering.append(parse_xml(
            f'<w:num {NSDECL} w:numId="{num_id}"><w:abstractNumId w:val="0"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'))
        return num_id

    def spacer(self, after=120):
        return self.raw(f'<w:p {NSDECL}><w:pPr><w:spacing w:after="{after}"/></w:pPr></w:p>')

    def page_break(self):
        return self.raw(f'<w:p {NSDECL}><w:r><w:br w:type="page"/></w:r></w:p>')

    # -- boxes -------------------------------------------------------------
    def takeaways(self, items, title="Key Takeaways"):
        box = copy.deepcopy(self.protos["takeaways"])
        tc = box.find(".//" + qn("w:tc"))
        for p in tc.findall(qn("w:p")):
            tc.remove(p)
        tc.append(parse_xml(p_xml(title, "BoxTitle")))
        for it in items:
            tc.append(parse_xml(p_xml(it, "BoxText",
                                      f'<w:numPr><w:ilvl w:val="0"/><w:numId w:val="{TAKEAWAY_NUM}"/></w:numPr>')))
        self.spacer(60)
        return self._add(box)

    def callout(self, title, lines):
        """Navy-bordered example box from the template (used for worked examples)."""
        box = copy.deepcopy(self.protos["callout"])
        tc = box.find(".//" + qn("w:tc"))
        for p in tc.findall(qn("w:p")):
            tc.remove(p)
        tc.append(parse_xml(p_xml(title, "BoxTitle", rpr="")))
        for ln in lines:
            self._count(ln)
            tc.append(parse_xml(p_xml(ln, "BoxText")))
        self._add(box)
        return self.spacer(120)

    # -- captions ----------------------------------------------------------
    def _caption(self, kind, title, center):
        counter = self.fig_no if kind == "Figure" else self.tab_no
        counter[self.chapter] = counter.get(self.chapter, 0) + 1
        n = counter[self.chapter]
        jc = '<w:jc w:val="center"/>' if center else ""
        keep = "<w:keepNext/>" if kind == "Table" else ""
        reset = " \\r 1" if kind in self.pending_reset else ""
        self.pending_reset.discard(kind)
        xml = (f'<w:p {NSDECL}><w:pPr><w:pStyle w:val="Caption"/>{keep}{jc}</w:pPr>'
               f'<w:r><w:t xml:space="preserve">{kind} {self.chapter}.</w:t></w:r>'
               '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
               f'<w:r><w:instrText xml:space="preserve">SEQ {kind} \\* ARABIC \\s 1{reset}</w:instrText></w:r>'
               '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
               f'<w:r><w:rPr><w:noProof/></w:rPr><w:t>{n}</w:t></w:r>'
               '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
               f'<w:r><w:t xml:space="preserve">: {escape(title)}</w:t></w:r></w:p>')
        self.raw(xml)
        return f"{kind} {self.chapter}.{n}"

    def source(self, text, center=False):
        jc = '<w:jc w:val="center"/>' if center else ""
        return self.raw(p_xml(text, "SourceNote", jc))

    def peek(self, kind):
        counter = self.fig_no if kind == "Figure" else self.tab_no
        return f"{kind} {self.chapter}.{counter.get(self.chapter, 0) + 1}"

    # -- figures -----------------------------------------------------------
    def figure(self, path, title, source, width_in=6.3):
        para = self.doc.add_paragraph()
        para.paragraph_format.keep_with_next = True
        run = para.add_run()
        run.add_picture(str(path), width=Inches(width_in))
        p = para._p
        ppr = p.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:spacing {NSDECL} w:before="120" w:after="40" w:line="240" w:lineRule="auto"/>'))
        ppr.append(parse_xml(f'<w:jc {NSDECL} w:val="center"/>'))
        self._add(p)
        label = self._caption("Figure", title, center=True)
        self.source(source, center=True)
        return label

    # -- tables ------------------------------------------------------------
    def table(self, title, header, rows, widths, source=None, bold_first_col=False, font_half_pts=None):
        assert abs(sum(widths) - 9360) <= 2, sum(widths)
        label = self._caption("Table", title, center=False)
        grid = "".join(f'<w:gridCol w:w="{w}"/>' for w in widths)
        sz = f'<w:sz w:val="{font_half_pts}"/><w:szCs w:val="{font_half_pts}"/>' if font_half_pts else ""

        def cell(text, w, head=False, shade=False, bold=False):
            fill = '<w:shd w:val="clear" w:color="auto" w:fill="1F3864"/>' if head else (
                '<w:shd w:val="clear" w:color="auto" w:fill="F5F7F7"/>' if shade else "")
                # alternating band, as in the template's sample rows
            va = '<w:vAlign w:val="center"/>' if head else ""
            style = "TableHead" if head else "TableText"
            rpr = sz + ('<w:b/><w:bCs/><w:color w:val="1F3864"/>' if bold and not head else "")
            paras = "".join(p_xml(line, style, "", rpr).replace(f"<w:p {NSDECL}>", "<w:p>")
                            for line in str(text).split("\n"))
            return (f'<w:tc><w:tcPr><w:tcW w:w="{w}" w:type="dxa"/>{fill}{CELL_MAR}{va}</w:tcPr>'
                    f"{paras}</w:tc>")

        trs = ['<w:tr><w:trPr><w:cantSplit/><w:tblHeader/></w:trPr>'
               + "".join(cell(h, w, head=True) for h, w in zip(header, widths)) + "</w:tr>"]
        for k, r in enumerate(rows):
            trs.append('<w:tr><w:trPr><w:cantSplit/></w:trPr>'
                       + "".join(cell(v, w, shade=(k % 2 == 1), bold=(bold_first_col and j == 0))
                                 for j, (v, w) in enumerate(zip(r, widths)))
                       + "</w:tr>")
        self.raw(f"<w:tbl {NSDECL}>{TBL_PR}<w:tblGrid>{grid}</w:tblGrid>{''.join(trs)}</w:tbl>")
        if source:
            self.source(source)
        else:
            self.spacer(120)
        return label


def set_runs_text(p_el, text, keep_first_rpr=True, plain=True):
    """Replace all runs of a paragraph with one run carrying the first run's formatting,
    minus the grey placeholder shading/italics when plain=True."""
    rs = p_el.findall(qn("w:r"))
    rpr = None
    if rs and keep_first_rpr and rs[0].find(qn("w:rPr")) is not None:
        rpr = copy.deepcopy(rs[0].find(qn("w:rPr")))
        if plain:
            for tag in ("w:i", "w:iCs", "w:shd"):
                for e in rpr.findall(qn(tag)):
                    rpr.remove(e)
            for e in rpr.findall(qn("w:color")):
                if e.get(qn("w:val")) in ("7F7F7F",):
                    rpr.remove(e)
    for r in rs:
        p_el.remove(r)
    for h in p_el.findall(qn("w:hyperlink")):
        p_el.remove(h)
    r = parse_xml(f'<w:r {NSDECL}><w:t xml:space="preserve">{escape(text)}</w:t></w:r>')
    if rpr is not None:
        r.insert(0, rpr)
    p_el.append(r)
    return p_el


def fill_cell(tc, text, bold_navy=False):
    ps = tc.findall(qn("w:p"))
    for extra in ps[1:]:
        tc.remove(extra)
    set_runs_text(ps[0], text)
    if bold_navy:
        r = ps[0].find(qn("w:r"))
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = parse_xml(f"<w:rPr {NSDECL}/>")
            r.insert(0, rpr)
        rpr.append(parse_xml(f"<w:b {NSDECL}/>"))
        rpr.append(parse_xml(f'<w:color {NSDECL} w:val="1F3864"/>'))
