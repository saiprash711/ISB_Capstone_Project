"""
Stunning PowerPoint Presentation Generator for EDA Dashboard
Creates a professional, visually appealing presentation showcasing the dashboard
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
import os

# Color Palette (matching dashboard theme)
COLOR_PRIMARY = RGBColor(2, 132, 199)  # Cyan Blue
COLOR_SECONDARY = RGBColor(99, 102, 241)  # Indigo
COLOR_ACCENT = RGBColor(0, 229, 255)  # Bright Cyan
COLOR_SUCCESS = RGBColor(16, 185, 129)  # Green
COLOR_WARNING = RGBColor(245, 158, 11)  # Amber
COLOR_DARK_BG = RGBColor(15, 23, 42)  # Dark Blue
COLOR_TEXT_PRIMARY = RGBColor(255, 255, 255)  # White
COLOR_TEXT_SECONDARY = RGBColor(203, 213, 225)  # Light Gray

def add_gradient_background(slide, color1, color2):
    """Add a gradient background to a slide"""
    background = slide.background
    fill = background.fill
    fill.gradient()
    fill.gradient_angle = 135
    fill.gradient_stops[0].color.rgb = color1
    fill.gradient_stops[1].color.rgb = color2

def add_title_with_icon(slide, title_text, icon_emoji, layout='title'):
    """Add a styled title with emoji icon"""
    if layout == 'title':
        title = slide.shapes.title
        title.text = f"{icon_emoji}  {title_text}"
        title.text_frame.paragraphs[0].font.size = Pt(44)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].font.color.rgb = COLOR_TEXT_PRIMARY
    else:
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
        title_frame = title_box.text_frame
        title_frame.text = f"{icon_emoji}  {title_text}"
        p = title_frame.paragraphs[0]
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY

def add_feature_card(slide, left, top, width, height, icon, title, description, color):
    """Add a feature card with icon, title, and description"""
    # Card background
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(25, 35, 55)
    card.line.color.rgb = RGBColor(50, 60, 90)
    card.line.width = Pt(1)
    card.shadow.inherit = False
    
    # Icon
    icon_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.15), Inches(0.5), Inches(0.5))
    icon_frame = icon_box.text_frame
    icon_frame.text = icon
    icon_frame.paragraphs[0].font.size = Pt(32)
    
    # Title
    title_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.6), width - Inches(0.3), Inches(0.4))
    title_frame = title_box.text_frame
    title_frame.text = title
    title_frame.paragraphs[0].font.size = Pt(16)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = color
    
    # Description
    desc_box = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(1.05), width - Inches(0.3), height - Inches(1.2))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    desc_frame.text = description
    p = desc_frame.paragraphs[0]
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_SECONDARY
    p.line_spacing = 1.3

def add_bullet_list(slide, left, top, width, height, items, title=None):
    """Add a bullet list"""
    text_box = slide.shapes.add_textbox(left, top, width, height)
    text_frame = text_box.text_frame
    text_frame.word_wrap = True
    
    if title:
        p = text_frame.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        p.space_after = Pt(12)
    
    for idx, item in enumerate(items):
        if idx == 0 and not title:
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(14)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(6)
        p.line_spacing = 1.2

def add_stat_card(slide, left, top, width, height, number, label, icon, color):
    """Add a statistics card"""
    # Card background
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(20, 28, 45)
    card.line.color.rgb = color
    card.line.width = Pt(2)
    
    # Icon
    icon_box = slide.shapes.add_textbox(left + Inches(0.1), top + Inches(0.1), width - Inches(0.2), Inches(0.4))
    icon_frame = icon_box.text_frame
    icon_frame.text = icon
    icon_frame.paragraphs[0].font.size = Pt(28)
    icon_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # Number
    num_box = slide.shapes.add_textbox(left + Inches(0.1), top + Inches(0.5), width - Inches(0.2), Inches(0.5))
    num_frame = num_box.text_frame
    num_frame.text = number
    p = num_frame.paragraphs[0]
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = color
    p.alignment = PP_ALIGN.CENTER
    
    # Label
    label_box = slide.shapes.add_textbox(left + Inches(0.1), top + Inches(1.05), width - Inches(0.2), Inches(0.35))
    label_frame = label_box.text_frame
    label_frame.text = label
    label_frame.word_wrap = True
    p = label_frame.paragraphs[0]
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_SECONDARY
    p.alignment = PP_ALIGN.CENTER

def create_presentation():
    """Create the complete presentation"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    print("🎨 Creating stunning PowerPoint presentation...")
    
    # ==================== SLIDE 1: Title Slide ====================
    print("  📄 Slide 1: Title slide...")
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    add_gradient_background(slide1, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    
    # Main title
    title_box = slide1.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = "⚡ Daikin AI Intelligence &\nDemand Forecasting Dashboard"
    p = title_frame.paragraphs[0]
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    subtitle_box = slide1.shapes.add_textbox(Inches(2), Inches(3.8), Inches(6), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Comprehensive EDA & Data Quality Assessment Platform\nISB Capstone Project"
    p = subtitle_frame.paragraphs[0]
    p.font.size = Pt(18)
    p.font.color.rgb = COLOR_TEXT_SECONDARY
    p.alignment = PP_ALIGN.CENTER
    
    # Accent line
    line = slide1.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(3.5), Inches(4.8), Inches(3), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    line.line.fill.background()
    
    # Stats at bottom
    stats = [
        ("📊", "7 Pages", Inches(1.5)),
        ("🎯", "82K+ Records", Inches(4)),
        ("⚡", "Real-time", Inches(6.5))
    ]
    for icon, text, left in stats:
        stat_box = slide1.shapes.add_textbox(left, Inches(6), Inches(2), Inches(0.5))
        stat_frame = stat_box.text_frame
        stat_frame.text = f"{icon}  {text}"
        p = stat_frame.paragraphs[0]
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        p.alignment = PP_ALIGN.CENTER
    
    # ==================== SLIDE 2: Dashboard Overview ====================
    print("  📄 Slide 2: Dashboard overview...")
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide2, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide2, "Dashboard Overview", "🎯", layout='content')
    
    # Key features
    features_items = [
        "Modern, gradient-based UI with purple-blue theme",
        "Interactive visualizations using Plotly",
        "Real-time data quality assessment",
        "Comprehensive statistical analysis tools",
        "Advanced demand forecasting with ML & statistical models",
        "Professional report generation (PDF, Excel, CSV)",
        "Responsive design for all screen sizes"
    ]
    add_bullet_list(slide2, Inches(0.7), Inches(1.3), Inches(4.5), Inches(5.5), features_items, "✨ Key Features")
    
    # Statistics
    add_stat_card(slide2, Inches(5.5), Inches(1.5), Inches(2), Inches(1.6), "7", "Dashboard\nPages", "📊", COLOR_ACCENT)
    add_stat_card(slide2, Inches(7.7), Inches(1.5), Inches(2), Inches(1.6), "23", "Data\nColumns", "📈", COLOR_SUCCESS)
    add_stat_card(slide2, Inches(5.5), Inches(3.3), Inches(2), Inches(1.6), "82K+", "Records\nAnalyzed", "🔢", COLOR_PRIMARY)
    add_stat_card(slide2, Inches(7.7), Inches(3.3), Inches(2), Inches(1.6), "100%", "Interactive", "⚡", COLOR_WARNING)
    
    # Tech stack
    tech_box = slide2.shapes.add_textbox(Inches(5.5), Inches(5.2), Inches(4.2), Inches(1.8))
    tech_frame = tech_box.text_frame
    tech_frame.word_wrap = True
    p = tech_frame.paragraphs[0]
    p.text = "🔧 Built With"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    
    p = tech_frame.add_paragraph()
    p.text = "Streamlit • Plotly • Pandas • NumPy\nStatsmodels • Scikit-learn • FPDF"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_SECONDARY
    p.space_before = Pt(8)
    
    # ==================== SLIDE 3: Seven Dashboard Pages ====================
    print("  📄 Slide 3: Seven pages overview...")
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide3, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide3, "Seven Comprehensive Pages", "📑", layout='content')
    
    pages_data = [
        ("🏠", "Overview", "Quick snapshot: metrics, trends,\nsegment distribution", COLOR_ACCENT),
        ("📊", "Data Quality", "Quality score, missing values,\ndata type analysis", COLOR_SUCCESS),
        ("🧹", "Data Cleaning", "Recommendations, duplicates,\nstandardization needs", COLOR_WARNING),
        ("📈", "Statistical Analysis", "Descriptive stats, correlations,\ndistributions", COLOR_PRIMARY)
    ]
    
    row = 0
    for idx, (icon, title, desc, color) in enumerate(pages_data):
        left = Inches(0.6) if idx % 2 == 0 else Inches(5.3)
        top = Inches(1.5 + (row * 1.45))
        add_feature_card(slide3, left, top, Inches(4.2), Inches(1.3), icon, title, desc, color)
        if idx % 2 == 1:
            row += 1
    
    pages_data2 = [
        ("🔍", "Detailed Exploration", "Column browser, search,\ndata preview", COLOR_SECONDARY),
        ("🎯", "Demand Forecasting", "SKU predictions, ML models,\nconfidence intervals", RGBColor(244, 63, 94)),
        ("📄", "Export Reports", "PDF, Excel, CSV exports\nwith full analysis", COLOR_SUCCESS)
    ]
    
    for idx, (icon, title, desc, color) in enumerate(pages_data2):
        left = Inches(0.6) if idx % 2 == 0 else Inches(5.3)
        top = Inches(4.4 + ((idx // 2) * 1.45))
        add_feature_card(slide3, left, top, Inches(4.2), Inches(1.3), icon, title, desc, color)
    
    # ==================== SLIDE 4: Overview Page ====================
    print("  📄 Slide 4: Overview page details...")
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide4, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide4, "🏠 Overview Page", "Dashboard Home", layout='content')
    
    # Add screenshot placeholder if exists
    screenshot_path = "screenshot_light_full.png"
    if os.path.exists(screenshot_path):
        slide4.shapes.add_picture(screenshot_path, Inches(0.5), Inches(1.2), width=Inches(5))
    
    # Features
    overview_features = [
        "📊 Key Metrics Cards",
        "  • Total records: 82,034",
        "  • Data completeness percentage",
        "  • Duplicate records count",
        "  • Missing values analysis",
        "",
        "📈 Visualizations",
        "  • Time series: Records over billing dates",
        "  • Gauge: Data completeness indicator",
        "  • Pie chart: Segment distribution",
        "  • Box plots: Tonnage analysis by segment"
    ]
    add_bullet_list(slide4, Inches(5.8), Inches(1.3), Inches(4), Inches(5.5), overview_features)
    
    # ==================== SLIDE 5: Data Quality & Cleaning ====================
    print("  📄 Slide 5: Data quality & cleaning...")
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide5, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide5, "Data Quality & Cleaning", "📊🧹", layout='content')
    
    # Left column - Data Quality
    quality_box = slide5.shapes.add_textbox(Inches(0.6), Inches(1.3), Inches(4.5), Inches(5.5))
    quality_frame = quality_box.text_frame
    quality_frame.word_wrap = True
    
    p = quality_frame.paragraphs[0]
    p.text = "📊 Data Quality Assessment"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p.space_after = Pt(12)
    
    quality_items = [
        "• Overall Quality Score (0-100)",
        "• Automated status indicators:",
        "  ✅ Excellent: >95% complete",
        "  ⚠️ Good: 75-95% complete",
        "  🔴 Poor: <75% complete",
        "",
        "• Missing values bar charts",
        "• Column-by-column analysis",
        "• Data type distribution",
        "• Issues counter and alerts"
    ]
    
    for item in quality_items:
        p = quality_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(4)
        p.line_spacing = 1.3
    
    # Right column - Data Cleaning
    cleaning_box = slide5.shapes.add_textbox(Inches(5.3), Inches(1.3), Inches(4.3), Inches(5.5))
    cleaning_frame = cleaning_box.text_frame
    cleaning_frame.word_wrap = True
    
    p = cleaning_frame.paragraphs[0]
    p.text = "🧹 Data Cleaning Studio"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_WARNING
    p.space_after = Pt(12)
    
    cleaning_items = [
        "• Actionable Recommendations:",
        "  ○ Duplicate removal strategies",
        "  ○ Missing value handling",
        "  ○ Column standardization",
        "  ○ Date format validation",
        "  ○ Categorical normalization",
        "",
        "• Visual Analysis:",
        "  ○ Unique vs duplicate pie chart",
        "  ○ Column standardization panel",
        "  ○ Top 5 values preview",
        "  ○ Before/after metrics"
    ]
    
    for item in cleaning_items:
        p = cleaning_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(4)
        p.line_spacing = 1.3
    
    # ==================== SLIDE 6: Statistical Analysis ====================
    print("  📄 Slide 6: Statistical analysis...")
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide6, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide6, "📈 Statistical Analysis", "Deep Insights", layout='content')
    
    # Features
    stat_items = [
        "📊 Descriptive Statistics Table",
        "  • Count, Mean, Std Dev, Min/Max",
        "  • Quartiles (25%, 50%, 75%)",
        "  • All numeric columns",
        "",
        "🔥 Correlation Heatmap",
        "  • Color-coded relationships (-1 to +1)",
        "  • Interactive hover for exact values",
        "  • Identify strong correlations",
        "",
        "📉 Distribution Analysis",
        "  • Interactive column selector",
        "  • Customizable histogram bins",
        "  • Box plot for outlier detection",
        "  • Statistical summary cards"
    ]
    add_bullet_list(slide6, Inches(0.7), Inches(1.3), Inches(4.8), Inches(5.5), stat_items)
    
    # Statistics cards
    add_stat_card(slide6, Inches(5.8), Inches(1.5), Inches(1.9), Inches(1.4), "Mean", "Average\nValues", "μ", COLOR_PRIMARY)
    add_stat_card(slide6, Inches(7.9), Inches(1.5), Inches(1.9), Inches(1.4), "σ", "Standard\nDeviation", "📊", COLOR_SUCCESS)
    add_stat_card(slide6, Inches(5.8), Inches(3.1), Inches(1.9), Inches(1.4), "Med", "Median\nValues", "⚖️", COLOR_ACCENT)
    add_stat_card(slide6, Inches(7.9), Inches(3.1), Inches(1.9), Inches(1.4), "R²", "Correlation\nStrength", "🔗", COLOR_WARNING)
    
    # Use cases
    usecase_box = slide6.shapes.add_textbox(Inches(5.8), Inches(4.7), Inches(4), Inches(2))
    usecase_frame = usecase_box.text_frame
    usecase_frame.word_wrap = True
    
    p = usecase_frame.paragraphs[0]
    p.text = "💡 Use Cases"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    
    p = usecase_frame.add_paragraph()
    p.text = "• Pattern identification & anomaly detection\n• Feature engineering for ML models\n• Research and exploratory analysis\n• Understanding data distributions"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_SECONDARY
    p.space_before = Pt(8)
    p.line_spacing = 1.4
    
    # ==================== SLIDE 7: Demand Forecasting (Star Feature) ====================
    print("  📄 Slide 7: Demand forecasting...")
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide7, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide7, "🎯 Demand Forecasting", "⭐ New Feature", layout='content')
    
    # Star badge
    star_badge = slide7.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(8.5), Inches(0.35), Inches(1.2), Inches(0.4)
    )
    star_badge.fill.solid()
    star_badge.fill.fore_color.rgb = RGBColor(244, 63, 94)
    star_badge.line.fill.background()
    
    badge_text = slide7.shapes.add_textbox(Inches(8.5), Inches(0.35), Inches(1.2), Inches(0.4))
    badge_frame = badge_text.text_frame
    badge_frame.text = "⭐ NEW"
    p = badge_frame.paragraphs[0]
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    
    # Models section
    models_box = slide7.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(4.5), Inches(3))
    models_frame = models_box.text_frame
    models_frame.word_wrap = True
    
    p = models_frame.paragraphs[0]
    p.text = "🤖 Available Models"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(244, 63, 94)
    p.space_after = Pt(12)
    
    models_content = [
        "📈 Statistical Models:",
        "  • ARIMA - Auto Regressive Integrated MA",
        "  • SARIMA - Seasonal ARIMA",
        "  • Exponential Smoothing",
        "",
        "🧠 Machine Learning Models:",
        "  • Random Forest Regressor",
        "  • Gradient Boosting Regressor",
        "  • Automatic feature engineering",
        "  • Lag features & rolling averages"
    ]
    
    for item in models_content:
        p = models_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(3)
        p.line_spacing = 1.4
    
    # Features section
    features_box = slide7.shapes.add_textbox(Inches(0.7), Inches(4.5), Inches(4.5), Inches(2.3))
    features_frame = features_box.text_frame
    features_frame.word_wrap = True
    
    p = features_frame.paragraphs[0]
    p.text = "✨ Key Features"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(10)
    
    features_content = [
        "• SKU-wise demand prediction (4-52 weeks)",
        "• Interactive forecast charts with confidence intervals",
        "• Performance metrics: MAE, RMSE, MAPE, Accuracy",
        "• Historical pattern analysis",
        "• CSV export for further analysis"
    ]
    
    for item in features_content:
        p = features_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(3)
        p.line_spacing = 1.4
    
    # Use cases
    usecase_box = slide7.shapes.add_textbox(Inches(5.5), Inches(1.3), Inches(4.2), Inches(5.5))
    usecase_frame = usecase_box.text_frame
    usecase_frame.word_wrap = True
    
    p = usecase_frame.paragraphs[0]
    p.text = "🎯 Business Use Cases"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p.space_after = Pt(12)
    
    usecases = [
        ("📦", "Inventory Planning", "Optimize stock levels based on\npredicted demand patterns"),
        ("🏭", "Production Capacity", "Schedule manufacturing runs\naccording to forecasts"),
        ("💰", "Budget Allocation", "Allocate resources by predicted\nSKU performance"),
        ("📈", "Sales Strategy", "Plan promotions during high\ndemand periods"),
        ("🚚", "Supply Chain", "Improve logistics and reduce\nstockouts/overstock")
    ]
    
    for icon, title, desc in usecases:
        p = usecase_frame.add_paragraph()
        p.text = f"{icon} {title}"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY
        p.space_before = Pt(10)
        
        p = usecase_frame.add_paragraph()
        p.text = desc
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(3)
        p.line_spacing = 1.2
    
    # ==================== SLIDE 8: Export & Reporting ====================
    print("  📄 Slide 8: Export & reporting...")
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide8, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide8, "📄 Export & Reporting", "Professional Output", layout='content')
    
    # Export formats
    formats_data = [
        ("📕", "PDF Reports", "Executive summary, quality\nassessment, recommendations", RGBColor(239, 68, 68)),
        ("📗", "Excel Export", "Full dataset + summary sheet\nfor further analysis", RGBColor(34, 197, 94)),
        ("📘", "CSV Export", "Raw data compatible with\nall analysis tools", RGBColor(59, 130, 246)),
        ("📄", "Text Summary", "Plain text reports for quick\nreference and email", RGBColor(168, 85, 247))
    ]
    
    for idx, (icon, title, desc, color) in enumerate(formats_data):
        row = idx // 2
        col = idx % 2
        left = Inches(0.7 + col * 4.8)
        top = Inches(1.5 + row * 1.7)
        add_feature_card(slide8, left, top, Inches(4.3), Inches(1.5), icon, title, desc, color)
    
    # Report contents
    contents_box = slide8.shapes.add_textbox(Inches(0.7), Inches(4.9), Inches(9), Inches(2.3))
    contents_frame = contents_box.text_frame
    contents_frame.word_wrap = True
    
    p = contents_frame.paragraphs[0]
    p.text = "📋 PDF Report Contents"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT
    p.space_after = Pt(10)
    
    p = contents_frame.add_paragraph()
    p.text = "Executive Summary  •  Data Quality Assessment  •  Missing Values Analysis  •  Data Type Distribution"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p.space_before = Pt(5)
    
    p = contents_frame.add_paragraph()
    p.text = "Cleaning Recommendations  •  Standardization Strategies  •  Monitoring Guidelines  •  Next Steps"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p.space_before = Pt(5)
    
    p = contents_frame.add_paragraph()
    p.text = "✅ Professional formatting  •  Automated generation  •  Ready for stakeholders"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_SUCCESS
    p.space_before = Pt(12)
    
    # ==================== SLIDE 9: Interactive Features ====================
    print("  📄 Slide 9: Interactive features...")
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide9, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide9, "⚡ Interactive Features", "Real-time Analysis", layout='content')
    
    # Chart interactions
    chart_box = slide9.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(4.5), Inches(2.8))
    chart_frame = chart_box.text_frame
    chart_frame.word_wrap = True
    
    p = chart_frame.paragraphs[0]
    p.text = "📊 Chart Interactions"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)
    
    chart_items = [
        "🖱️ Hover: View exact values",
        "🔍 Zoom: Click and drag to zoom in",
        "↔️ Pan: Shift + drag to move",
        "💾 Download: Save charts as PNG",
        "🔄 Reset: Double-click to reset view",
        "🎨 Interactive legends: Toggle series"
    ]
    
    for item in chart_items:
        p = chart_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(5)
        p.line_spacing = 1.3
    
    # Data filtering
    filter_box = slide9.shapes.add_textbox(Inches(0.7), Inches(4.3), Inches(4.5), Inches(2.5))
    filter_frame = filter_box.text_frame
    filter_frame.word_wrap = True
    
    p = filter_frame.paragraphs[0]
    p.text = "🔍 Data Filtering & Search"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p.space_after = Pt(12)
    
    filter_items = [
        "• Multi-select column explorer",
        "• Search within specific columns",
        "• Time range analysis",
        "• Individual column deep-dive",
        "• Value counts and distributions",
        "• First 100 rows preview"
    ]
    
    for item in filter_items:
        p = filter_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(5)
        p.line_spacing = 1.3
    
    # Visualizations types
    viz_box = slide9.shapes.add_textbox(Inches(5.5), Inches(1.3), Inches(4.2), Inches(5.5))
    viz_frame = viz_box.text_frame
    viz_frame.word_wrap = True
    
    p = viz_frame.paragraphs[0]
    p.text = "📈 Supported Visualizations"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_WARNING
    p.space_after = Pt(12)
    
    viz_types = [
        ("📉", "Line Charts", "Time series trends and patterns"),
        ("📊", "Bar Charts", "Categorical comparisons"),
        ("🥧", "Pie Charts", "Proportions and distributions"),
        ("📊", "Histograms", "Value distributions"),
        ("📦", "Box Plots", "Outliers and quartiles"),
        ("🔥", "Heatmaps", "Correlation matrices"),
        ("🎯", "Gauge Charts", "Single metric visualization"),
        ("🔵", "Scatter Plots", "Relationship analysis")
    ]
    
    for icon, title, desc in viz_types:
        p = viz_frame.add_paragraph()
        p.text = f"{icon} {title}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY
        p.space_before = Pt(6)
        
        p = viz_frame.add_paragraph()
        p.text = f"   {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(2)
    
    # ==================== SLIDE 10: Technical Architecture ====================
    print("  📄 Slide 10: Technical architecture...")
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide10, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide10, "🔧 Technical Architecture", "Under the Hood", layout='content')
    
    # Tech stack
    add_stat_card(slide10, Inches(0.7), Inches(1.4), Inches(2.2), Inches(1.4), "Python", "Core\nLanguage", "🐍", COLOR_PRIMARY)
    add_stat_card(slide10, Inches(3.1), Inches(1.4), Inches(2.2), Inches(1.4), "Streamlit", "Web\nFramework", "⚡", COLOR_ACCENT)
    add_stat_card(slide10, Inches(5.5), Inches(1.4), Inches(2.2), Inches(1.4), "Plotly", "Interactive\nCharts", "📊", COLOR_SUCCESS)
    add_stat_card(slide10, Inches(7.9), Inches(1.4), Inches(1.9), Inches(1.4), "Pandas", "Data\nAnalysis", "🐼", COLOR_WARNING)
    
    # Libraries
    libs_box = slide10.shapes.add_textbox(Inches(0.7), Inches(3), Inches(4.5), Inches(3.8))
    libs_frame = libs_box.text_frame
    libs_frame.word_wrap = True
    
    p = libs_frame.paragraphs[0]
    p.text = "📚 Core Libraries"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(12)
    
    libs_items = [
        "🌐 Streamlit - Web framework",
        "📊 Plotly - Interactive visualizations",
        "🐼 Pandas - Data manipulation",
        "🔢 NumPy - Numerical computing",
        "📈 Statsmodels - Time series models",
        "🤖 Scikit-learn - Machine learning",
        "📄 FPDF - PDF generation",
        "📗 OpenPyXL - Excel handling"
    ]
    
    for item in libs_items:
        p = libs_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(5)
        p.line_spacing = 1.3
    
    # Features
    features_box = slide10.shapes.add_textbox(Inches(5.5), Inches(3), Inches(4.2), Inches(3.8))
    features_frame = features_box.text_frame
    features_frame.word_wrap = True
    
    p = features_frame.paragraphs[0]
    p.text = "⚡ Performance & Optimization"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p.space_after = Pt(12)
    
    perf_items = [
        "• Cached data loading for speed",
        "• Efficient pandas operations",
        "• Lazy loading of visualizations",
        "• Handles 82K+ records smoothly",
        "• Responsive UI with smooth animations",
        "• Graceful handling of missing data",
        "• Optimized memory usage",
        "• Fast duplicate detection algorithms"
    ]
    
    for item in perf_items:
        p = features_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(5)
        p.line_spacing = 1.3
    
    # ==================== SLIDE 11: Use Cases & Benefits ====================
    print("  📄 Slide 11: Use cases & benefits...")
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide11, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide11, "💡 Use Cases & Benefits", "Who Benefits?", layout='content')
    
    usecases_data = [
        ("👨‍💼", "Data Analysts", [
            "• Quick EDA before modeling",
            "• Pattern identification",
            "• Outlier detection",
            "• Distribution analysis"
        ], COLOR_PRIMARY),
        ("👨‍💻", "Data Engineers", [
            "• Pipeline monitoring",
            "• Data validation",
            "• Quality tracking",
            "• Cleaning recommendations"
        ], COLOR_SUCCESS),
        ("👨‍💼", "Managers", [
            "• Executive summaries",
            "• Quality reports",
            "• Progress tracking",
            "• Stakeholder presentations"
        ], COLOR_WARNING),
        ("👨‍🔬", "Data Scientists", [
            "• Feature understanding",
            "• Correlation discovery",
            "• Statistical insights",
            "• Demand forecasting"
        ], COLOR_ACCENT)
    ]
    
    for idx, (icon, title, items, color) in enumerate(usecases_data):
        row = idx // 2
        col = idx % 2
        left = Inches(0.7 + col * 4.8)
        top = Inches(1.4 + row * 2.6)
        
        # Card
        card = slide11.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            left, top, Inches(4.3), Inches(2.3)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(20, 28, 45)
        card.line.color.rgb = color
        card.line.width = Pt(2)
        
        # Icon & Title
        title_box = slide11.shapes.add_textbox(left + Inches(0.15), top + Inches(0.15), Inches(4), Inches(0.4))
        title_frame = title_box.text_frame
        title_frame.text = f"{icon}  {title}"
        p = title_frame.paragraphs[0]
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = color
        
        # Items
        items_box = slide11.shapes.add_textbox(left + Inches(0.15), top + Inches(0.65), Inches(4), Inches(1.5))
        items_frame = items_box.text_frame
        items_frame.word_wrap = True
        
        for item in items:
            p = items_frame.add_paragraph() if items.index(item) > 0 else items_frame.paragraphs[0]
            p.text = item
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_TEXT_SECONDARY
            p.space_before = Pt(3)
            p.line_spacing = 1.2
    
    # ==================== SLIDE 12: Getting Started ====================
    print("  📄 Slide 12: Getting started...")
    slide12 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide12, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    add_title_with_icon(slide12, "🚀 Getting Started", "Quick Setup", layout='content')
    
    # Installation steps
    install_box = slide12.shapes.add_textbox(Inches(0.7), Inches(1.3), Inches(4.5), Inches(5.5))
    install_frame = install_box.text_frame
    install_frame.word_wrap = True
    
    p = install_frame.paragraphs[0]
    p.text = "📦 Installation Steps"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_PRIMARY
    p.space_after = Pt(16)
    
    steps = [
        ("1️⃣", "Prerequisites", "Python 3.8+ and pip installed"),
        ("2️⃣", "Install Dependencies", "pip install -r requirements.txt"),
        ("3️⃣", "Prepare Data", "Place Filtered.xlsx in project folder"),
        ("4️⃣", "Run Dashboard", "streamlit run dashboard.py"),
        ("5️⃣", "Open Browser", "Navigate to http://localhost:8501")
    ]
    
    for num, title, desc in steps:
        p = install_frame.add_paragraph()
        p.text = f"{num} {title}"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = COLOR_TEXT_PRIMARY
        p.space_before = Pt(14)
        
        p = install_frame.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(4)
    
    # Best practices
    best_box = slide12.shapes.add_textbox(Inches(5.5), Inches(1.3), Inches(4.2), Inches(5.5))
    best_frame = best_box.text_frame
    best_frame.word_wrap = True
    
    p = best_frame.paragraphs[0]
    p.text = "✨ Best Practices"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = COLOR_SUCCESS
    p.space_after = Pt(16)
    
    practices = [
        "🏠 Start with Overview page",
        "📊 Check Data Quality for issues",
        "🧹 Review Cleaning recommendations",
        "📈 Dive into Statistical Analysis",
        "🔍 Use Detailed Exploration for specifics",
        "🎯 Try Demand Forecasting for SKUs",
        "📄 Export Reports for documentation",
        "⚡ Use search for quick record lookup",
        "📌 Monitor trends over time",
        "🔄 Run weekly for quality tracking"
    ]
    
    for practice in practices:
        p = best_frame.add_paragraph()
        p.text = practice
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_TEXT_SECONDARY
        p.space_before = Pt(7)
        p.line_spacing = 1.3
    
    # ==================== SLIDE 13: Thank You / Contact ====================
    print("  📄 Slide 13: Thank you slide...")
    slide13 = prs.slides.add_slide(prs.slide_layouts[6])
    add_gradient_background(slide13, RGBColor(10, 14, 26), RGBColor(15, 23, 42))
    
    # Thank you title
    thanks_box = slide13.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
    thanks_frame = thanks_box.text_frame
    thanks_frame.text = "Thank You! 🎉"
    p = thanks_frame.paragraphs[0]
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_PRIMARY
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    sub_box = slide13.shapes.add_textbox(Inches(2), Inches(3.7), Inches(6), Inches(0.6))
    sub_frame = sub_box.text_frame
    sub_frame.text = "Daikin AI Intelligence & Demand Forecasting Dashboard"
    p = sub_frame.paragraphs[0]
    p.font.size = Pt(18)
    p.font.color.rgb = COLOR_TEXT_SECONDARY
    p.alignment = PP_ALIGN.CENTER
    
    # ISB Capstone
    isb_box = slide13.shapes.add_textbox(Inches(2), Inches(4.4), Inches(6), Inches(0.5))
    isb_frame = isb_box.text_frame
    isb_frame.text = "ISB Capstone Project"
    p = isb_frame.paragraphs[0]
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_ACCENT
    p.alignment = PP_ALIGN.CENTER
    
    # Features summary
    summary_box = slide13.shapes.add_textbox(Inches(2), Inches(5.3), Inches(6), Inches(1.2))
    summary_frame = summary_box.text_frame
    summary_frame.text = "7 Pages  •  82K+ Records  •  Real-time Analysis  •  ML Forecasting\nInteractive Visualizations  •  Professional Reports"
    p = summary_frame.paragraphs[0]
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_SECONDARY
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.5
    
    # Bottom accent line
    line = slide13.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(3), Inches(6.8), Inches(4), Inches(0.05)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_ACCENT
    line.line.fill.background()
    
    # Save presentation
    output_file = "Daikin_EDA_Dashboard_Presentation.pptx"
    prs.save(output_file)
    print(f"\n✅ Presentation created successfully!")
    print(f"📁 File: {output_file}")
    print(f"📊 Total slides: {len(prs.slides)}")
    print(f"\n🎨 Features:")
    print(f"   • Stunning gradient backgrounds")
    print(f"   • Professional card-based layouts")
    print(f"   • Color-coded sections")
    print(f"   • Interactive feature cards")
    print(f"   • Statistics visualizations")
    print(f"   • Comprehensive coverage of all 7 pages")
    print(f"\n💡 The presentation is ready to wow your audience!")

if __name__ == "__main__":
    create_presentation()
